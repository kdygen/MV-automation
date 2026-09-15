"""Turning an uploaded file into plain text. No model, no network, no OCR.

Three formats, one contract: given bytes and a filename, return normalized text plus
enough structure to say which page a passage came from — or raise
:class:`ExtractionError` carrying a sentence the document's owner can act on.

**Why the failure messages matter so much here.** The owner uploads a file and walks
away. If ingestion fails, the only thing they will ever see is ``failure_reason``, so it
has to name the actual problem and the fix ("this looks like a scan — upload a
text-based PDF"), not "processing failed". Every raise below is written to be read by a
moving-company owner, not by us.

**Scanned PDFs fail deliberately.** A page of images extracts as a handful of stray
characters. Indexing that would produce a document that looks ready, answers nothing,
and quietly poisons retrieval with fragments. OCR is a future decision with its own
accuracy and cost questions; until then the honest outcome is a clear failure.

**Uploaded files are data.** Nothing extracted here is ever treated as an instruction.
Markup is stripped during normalization and the text only ever reaches the model as
retrieved evidence, through the same tool contract as manual entries.
"""

from __future__ import annotations

import enum
import io
import re
from dataclasses import dataclass
from typing import TYPE_CHECKING

from app.core.logging import get_logger
from app.knowledge.chunking import normalize

if TYPE_CHECKING:  # pragma: no cover - import cost paid only by the type checker
    from docx.table import Table
    from docx.text.paragraph import Paragraph

logger = get_logger(__name__)


class DocumentKind(enum.StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    TEXT = "text"


#: Canonical MIME per kind. Stored on the row rather than the browser's guess, which is
#: attacker-controlled and frequently wrong even when it is not.
MIME_TYPES: dict[DocumentKind, str] = {
    DocumentKind.PDF: "application/pdf",
    DocumentKind.DOCX: (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    DocumentKind.TEXT: "text/plain",
}

TEXT_EXTENSIONS = frozenset({".txt", ".text", ".md", ".markdown"})

#: Below this, there is nothing worth indexing whatever the file claimed to be.
MIN_TEXT_CHARS = 80

#: A text-based PDF page carries hundreds of characters. A scanned one carries a few
#: stray marks from the OCR-less text layer. 40 sits far enough below real prose to
#: avoid false accusations and far enough above noise to catch scans reliably.
MIN_CHARS_PER_PAGE = 40

#: A heading in a Word document is only re-emitted as a Markdown heading when it is
#: short enough to actually be one; a 300-character "heading" is a paragraph with the
#: wrong style applied, and marking it up would corrupt the chunk text.
MAX_DOCX_HEADING_CHARS = 80

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"
_NULL_BYTE = b"\x00"
_EXTENSION_RE = re.compile(r"(\.[A-Za-z0-9]{1,10})$")


class ExtractionError(Exception):
    """Extraction failed. ``args[0]`` is shown to the owner verbatim."""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


@dataclass(frozen=True)
class ExtractedDocument:
    """Normalized text, plus where its pages start.

    ``page_starts`` maps a page number to its offset in :attr:`text`, which is what lets
    a chunk report the pages it came from without the chunker knowing anything about
    pages. Empty for formats that have no fixed pagination — a ``.docx`` has pages only
    once a word processor lays it out, and inventing numbers would be worse than none.
    """

    text: str
    kind: DocumentKind
    page_count: int | None
    page_starts: tuple[tuple[int, int], ...] = ()


def file_extension(filename: str) -> str:
    match = _EXTENSION_RE.search(filename or "")
    return match.group(1).lower() if match else ""


def detect_kind(filename: str, data: bytes) -> DocumentKind:
    """Decide what a file actually is, preferring its content over its name.

    Magic bytes win because the filename and the browser's content type are both
    supplied by the client. The extension only decides between formats that share a
    container — every ``.docx`` is a zip, but not every zip is a ``.docx``.
    """
    extension = file_extension(filename)

    if data.startswith(_PDF_MAGIC):
        return DocumentKind.PDF
    if data.startswith(_ZIP_MAGIC):
        if extension == ".docx":
            return DocumentKind.DOCX
        raise ExtractionError(
            "This looks like a zip archive. Upload the document itself, not a folder."
        )
    if extension == ".pdf":
        raise ExtractionError("This file is named .pdf but is not a PDF. Re-export it.")
    if extension == ".doc":
        raise ExtractionError(
            "Older .doc files are not supported. Save it as .docx or PDF and try again."
        )
    if extension in TEXT_EXTENSIONS:
        return DocumentKind.TEXT
    raise ExtractionError(
        "Unsupported file type. Upload a PDF, a Word document (.docx), or a text file."
    )


def extract(filename: str, data: bytes) -> ExtractedDocument:
    """Extract normalized text, or raise :class:`ExtractionError` with a usable reason."""
    if not data:
        raise ExtractionError("This file is empty.")

    kind = detect_kind(filename, data)
    extracted = _EXTRACTORS[kind](data)

    if len(extracted.text) < MIN_TEXT_CHARS:
        raise ExtractionError(_too_little_text(extracted))
    return extracted


def _too_little_text(extracted: ExtractedDocument) -> str:
    if extracted.kind is DocumentKind.PDF and (extracted.page_count or 0) > 0:
        return (
            "This looks like a scanned document — its pages hold images rather than "
            "text. Upload a text-based PDF, or add the policy as a knowledge entry."
        )
    return "No readable text was found in this file."


def _extract_pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # An empty user password is common on "protected" PDFs and decrypts cleanly;
            # anything else genuinely needs a password we will never ask the owner for.
            try:
                unlocked = bool(reader.decrypt(""))
            except Exception:  # noqa: BLE001 - any failure here means "still locked"
                unlocked = False
            if not unlocked:
                raise ExtractionError(
                    "This PDF is password-protected. Remove the password and upload it "
                    "again."
                )
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ExtractionError:
        raise
    except (PyPdfError, OSError, ValueError, KeyError, TypeError) as exc:
        logger.info("PDF extraction failed: %s", type(exc).__name__)
        raise ExtractionError(
            "This PDF could not be read. It may be damaged — try re-exporting it."
        ) from exc

    text, page_starts = _assemble_pages(pages)
    page_count = len(pages)

    # Separate from the global minimum: a 40-page scan can still scrape together 80
    # characters of noise, which would otherwise pass as a very short document.
    if page_count and len(text) < MIN_CHARS_PER_PAGE * page_count:
        raise ExtractionError(
            "This looks like a scanned document — its pages hold images rather than "
            "text. Upload a text-based PDF, or add the policy as a knowledge entry."
        )
    return ExtractedDocument(
        text=text, kind=DocumentKind.PDF, page_count=page_count, page_starts=page_starts
    )


def _assemble_pages(pages: list[str]) -> tuple[str, tuple[tuple[int, int], ...]]:
    """Join pages into one normalized string, remembering where each one begins.

    Each page is normalized *before* joining so the offsets recorded here stay valid
    when the chunker normalizes the joined text again — ``normalize`` is idempotent, so
    the string it produces is byte-identical to the one measured here.
    """
    parts: list[str] = []
    starts: list[tuple[int, int]] = []
    offset = 0
    for number, raw in enumerate(pages, start=1):
        cleaned = normalize(raw)
        if not cleaned:
            continue
        if parts:
            offset += 2  # the "\n\n" joined in below
        starts.append((number, offset))
        parts.append(cleaned)
        offset += len(cleaned)
    return "\n\n".join(parts), tuple(starts)


def _extract_docx(data: bytes) -> ExtractedDocument:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises a wide, untyped family
        logger.info("DOCX extraction failed: %s", type(exc).__name__)
        raise ExtractionError(
            "This Word document could not be read. It may be damaged or password-"
            "protected — try re-saving it as .docx or PDF."
        ) from exc

    blocks: list[str] = []
    # Walking the body in document order keeps a table with its introducing sentence.
    # Reading ``paragraphs`` then ``tables`` would move every table to the end, which
    # separates rate tables from the clause that explains them.
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(_docx_paragraph(Paragraph(child, document)))
        elif child.tag == qn("w:tbl"):
            blocks.append(_docx_table(Table(child, document)))

    text = normalize("\n\n".join(block for block in blocks if block))
    return ExtractedDocument(text=text, kind=DocumentKind.DOCX, page_count=None)


def _docx_paragraph(paragraph: Paragraph) -> str:
    """Return the paragraph's text, marking real headings so the chunker sees them."""
    text = (paragraph.text or "").strip()
    if not text:
        return ""
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    if style.startswith(("Heading", "Title")) and len(text) <= MAX_DOCX_HEADING_CHARS:
        return f"# {text}"
    return text


def _docx_table(table: Table) -> str:
    """Flatten a table to one line per row.

    Pipe-separated rather than reconstructed as a grid: the row is what carries meaning
    ("Piano | $250 flat"), and the lexical index matches words, not columns.
    """
    rows: list[str] = []
    for row in table.rows:
        cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
        # A merged cell repeats its text across the span; collapse consecutive repeats.
        deduped = [c for i, c in enumerate(cells) if c and (i == 0 or c != cells[i - 1])]
        if deduped:
            rows.append(" | ".join(deduped))
    return "\n".join(rows)


def _extract_text(data: bytes) -> ExtractedDocument:
    if _NULL_BYTE in data[:8192]:
        raise ExtractionError(
            "This file is not readable text. Upload a PDF, a .docx, or a plain text file."
        )
    try:
        decoded = data.decode("utf-8")
    except UnicodeDecodeError:
        # Latin-1 decodes any byte string, so this is a last resort rather than a guess
        # at the real encoding — it keeps a Windows-authored policy readable instead of
        # rejecting it over a single smart quote.
        decoded = data.decode("latin-1", errors="replace")
    return ExtractedDocument(
        text=normalize(decoded), kind=DocumentKind.TEXT, page_count=None
    )


_EXTRACTORS = {
    DocumentKind.PDF: _extract_pdf,
    DocumentKind.DOCX: _extract_docx,
    DocumentKind.TEXT: _extract_text,
}


# --------------------------------------------------------------------------- pages


def page_at(page_starts: tuple[tuple[int, int], ...], offset: int) -> int | None:
    """Which page contains ``offset``, or ``None`` when the format has no pages."""
    found: int | None = None
    for number, start in page_starts:
        if start <= offset:
            found = number
        else:
            break
    return found


def attribute_pages(
    text: str, contents: list[str], page_starts: tuple[tuple[int, int], ...]
) -> list[tuple[int | None, int | None]]:
    """Map each chunk back to the page range it came from. Best effort by design.

    Chunk content is the extracted text plus a heading prefix that does not appear in the
    source, so the span is measured from the paragraphs that *are* found rather than from
    ``len(content)`` — otherwise the prefix pushes every chunk's end past the page it
    really occupies. The cursor only moves forward, which keeps a phrase repeated across
    a document from dragging an early chunk to a late page.

    A chunk that cannot be located reports ``(None, None)``. That is the right answer:
    an approximate page number on a legal document is worse than no page number.
    """
    if not page_starts:
        return [(None, None)] * len(contents)

    spans: list[tuple[int | None, int | None]] = []
    cursor = 0
    for content in contents:
        start = -1
        end = -1
        for paragraph in content.split("\n\n"):
            probe = paragraph.strip()[:120]
            if not probe:
                continue
            found = text.find(probe, cursor)
            if found < 0:
                continue
            if start < 0:
                start = found
            end = found + len(probe)
        if start < 0:
            spans.append((None, None))
            continue
        cursor = start
        spans.append((page_at(page_starts, start), page_at(page_starts, max(end - 1, start))))
    return spans
