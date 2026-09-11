"""Reading CSV and XLSX uploads into one uniform row shape.

Both formats normalize to ``list[dict[str, object]]`` keyed by the original header text,
so every downstream layer — suggestion, validation, fingerprinting — is written once and
is format-agnostic. XLSX values arrive already typed by Excel (dates as ``datetime``,
numbers as ``float``), which the parsers accept alongside strings.

Size limits are enforced here rather than at the route because the expensive part is
materializing rows, not receiving bytes.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from app.core.errors import ValidationError

#: A year of moves for a busy company is a few thousand rows. Ten thousand is generous
#: and bounds the memory a single upload can consume.
MAX_ROWS = 10_000
MAX_COLUMNS = 200


@dataclass(frozen=True)
class Sheet:
    """A parsed upload: original headers, in order, plus its rows."""

    headers: tuple[str, ...]
    rows: tuple[dict[str, object], ...]

    def column(self, header: str) -> list[object]:
        """Every value in one column — what the ambiguity detectors scan."""
        return [row.get(header) for row in self.rows]


def read_upload(filename: str, content: bytes) -> Sheet:
    """Parse an uploaded file by extension.

    :raises app.core.errors.ValidationError: unreadable, empty, or oversized.
    """
    lowered = (filename or "").lower()
    if lowered.endswith(".xlsx"):
        return _read_xlsx(content)
    if lowered.endswith((".csv", ".txt")):
        return _read_csv(content)
    raise ValidationError("Upload a .csv or .xlsx file")


def _finish(headers: list[str], rows: list[dict[str, object]]) -> Sheet:
    clean_headers = [h.strip() for h in headers if h and h.strip()]
    if not clean_headers:
        raise ValidationError("The file has no column headers")
    if len(clean_headers) > MAX_COLUMNS:
        raise ValidationError(f"Files are limited to {MAX_COLUMNS} columns")
    if not rows:
        raise ValidationError("The file has headers but no rows")
    if len(rows) > MAX_ROWS:
        raise ValidationError(f"Files are limited to {MAX_ROWS:,} rows per upload")
    return Sheet(headers=tuple(clean_headers), rows=tuple(rows))


def _read_csv(content: bytes) -> Sheet:
    try:
        text = content.decode("utf-8-sig")  # tolerate Excel's BOM
    except UnicodeDecodeError:
        try:
            text = content.decode("latin-1")  # common in older exports
        except UnicodeDecodeError as exc:  # pragma: no cover - latin-1 accepts all bytes
            raise ValidationError("The CSV could not be decoded") from exc

    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # single-column files defeat the sniffer; comma is fine

    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if reader.fieldnames is None:
        raise ValidationError("The CSV file is empty")

    rows: list[dict[str, object]] = []
    for raw in reader:
        row = {(k or "").strip(): v for k, v in raw.items() if k and k.strip()}
        if any(str(v).strip() for v in row.values() if v is not None):
            rows.append(row)  # skip blank trailing lines
        if len(rows) > MAX_ROWS:
            break
    return _finish([h.strip() for h in reader.fieldnames if h], rows)


def _read_xlsx(content: bytes) -> Sheet:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ValidationError("XLSX support is unavailable on this server") from exc

    try:
        # read_only streams rows instead of building the whole object graph; data_only
        # takes the cached value of a formula rather than the formula text.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("The spreadsheet could not be read") from exc

    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration as exc:
            raise ValidationError("The spreadsheet is empty") from exc

        headers = [str(h).strip() if h is not None else "" for h in header_row]
        rows: list[dict[str, object]] = []
        for values in iterator:
            if not any(v is not None and str(v).strip() for v in values):
                continue
            rows.append(
                {h: v for h, v in zip(headers, values, strict=False) if h}
            )
            if len(rows) > MAX_ROWS:
                break
        return _finish(headers, rows)
    finally:
        workbook.close()
